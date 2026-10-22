import io
import unittest
from datetime import date, timedelta
from unittest.mock import Mock, patch
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Color
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from gcalveslo import (parse_day, event_key, sync, EXCLUDED, is_excluded,
    load_events, workbook_events, is_managed_event, canonical_title, cancelled_row)

class Tests(unittest.TestCase):
    def test_dates(self):
        self.assertEqual(parse_day('29.5.-1.6.',2027,5),(date(2027,5,29),date(2027,6,1)))
        self.assertEqual(parse_day('25.-26.1.',2027,1),(date(2027,1,25),date(2027,1,26)))
        self.assertEqual(parse_day('1.-3.',2027,9),(date(2027,9,1),date(2027,9,3)))
    def test_exclusions(self):
        self.assertTrue(EXCLUDED.search('Mistrovství světa ve veslování'))
        self.assertTrue(EXCLUDED.search('Světový pohár'))
        self.assertFalse(EXCLUDED.search('Pohár ČR'))
        for name in (
            'World Rowing Indoor Championships',
            'World Rowing Coastal Championships',
            'World Rowing Beach Sprint Championships',
            'World Rowing Under 23 Championships',
            'World Rowing Masters Championships',
            'World Rowing Junior Championships',
            'World Championships',
            'European Rowing Coastal Championships',
            'European Indoor Rowing Championships',
            'Indoor Championships',
            'Coastal Rowing Championships',
            'World Rowing Cup II',
            'MS ve veslování', 'ME ve veslování',
            'WRICH 2027', 'WRCH', 'ERICH', 'ERCH', 'WRBSCH',
            'World Rowing Para Rowing Championships',
            'Para Rowing International Regatta', 'Para-veslování',
            'World Rowing Beach Sprint Series I.kolo',
            'World Rowing Beach Sprint Finals',
            'YOG Dakar 2026 - coastal rowing',
            'Coupe de la Jeunesse 2027', '4th European Games',
        ):
            with self.subTest(name=name):
                self.assertTrue(EXCLUDED.search(name), name)
        for name in ('Mistrovství ČR', 'Národní mistrovství', 'Pražské primátorky', 'International Regatta', 'Kontrolní závody juniorů', 'Kontrolní závody seniorů', 'Kontrolní závody seniorů A', 'Český pohár v beach sprintu'):
            with self.subTest(name=name):
                self.assertFalse(EXCLUDED.search(name), name)
    def test_excluded_items_never_reach_google_even_if_supplied_directly(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {'items': []}
        base = dict(start=date(2027, 6, 19), end=date(2027, 6, 21),
                    description='', location='', international=True, source='https://veslo.cz/a')
        bad_titles = (
            'World Rowing Beach Sprint Finals',
            'World Rowing Beach Sprint Series I.kolo',
            'World Rowing Indoor Championships',
            'WRICH 2027',
            'World Rowing Para Rowing Championships',
            '4th European Games',
            'Coupe de la Jeunesse 2027',
            'YOG Dakar 2026 - coastal rowing',
        )
        good = dict(base, summary='Kontrolní závody seniorů A', international=False)
        events = [dict(base, summary=name) for name in bad_titles] + [good]
        self.assertEqual(sync(service, 'calendar', events, date(2027, 1, 1),
                              date(2028, 1, 1), '9'), (1, len(bad_titles), 0))
        calls = service.events.return_value.insert.call_args_list
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0].kwargs['body']['summary'], good['summary'])

    def test_same_name_day_even_if_end_differs(self):
        self.assertEqual(event_key(dict(summary='  ČESKÁ  REGATA ', start=date(2027,5,2), end=date(2027,5,3))),
                         event_key(dict(summary='česká regata', start=date(2027,5,2), end=date(2027,5,4))))

    def test_senior_control_race_alias(self):
        a = dict(summary='Kontrolní závody seniorů A', start=date(2027, 6, 5))
        b = dict(summary='Kontrolní závody seniorů', start=date(2027, 6, 5))
        self.assertEqual(event_key(a), event_key(b))
        self.assertNotEqual(event_key(a), event_key(dict(b, start=date(2027, 6, 6))))

    def test_existing_senior_control_race_alias(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {
            'items': [{'summary': 'Kontrolní závody seniorů A', 'start': {'date': '2027-06-05'}}]}
        event = dict(summary='Kontrolní závody seniorů', start=date(2027, 6, 5),
                     end=date(2027, 6, 7), description='', location='', international=False,
                     source='https://veslo.cz/a')
        self.assertEqual(sync(service, 'calendar', [event], date(2027, 1, 1), date(2028, 1, 1), '9'), (0, 1, 0))
        service.events.return_value.insert.assert_not_called()

    def test_exact_existing_duplicate_is_skipped(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {
            'items': [
                {'summary': 'Kontrolní závody seniorů A', 'start': {'date': '2027-06-05'}},
                {'summary': 'Kontrolní závody seniorů A', 'start': {'date': '2027-06-05'}},
            ]}
        event = dict(summary='Kontrolní závody seniorů A', start=date(2027, 6, 5),
                     end=date(2027, 6, 7), description='', location='', international=False,
                     source='https://veslo.cz/a')
        with self.assertLogs('gcalveslo', level='WARNING') as messages:
            self.assertEqual(sync(service, 'calendar', [event], date(2027, 1, 1), date(2028, 1, 1), '9'), (0, 1, 0))
        self.assertTrue(any('Already duplicated' in msg for msg in messages.output))
        service.events.return_value.insert.assert_not_called()
        service.events.return_value.delete.assert_not_called()

    def test_identical_races_in_same_import_insert_only_once(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {'items': []}
        event = dict(summary='Kontrolní závody seniorů A', start=date(2027, 6, 5),
                     end=date(2027, 6, 7), description='', location='', international=False,
                     source='https://veslo.cz/a')
        self.assertEqual(sync(service, 'calendar', [event, dict(event)], date(2027, 1, 1),
                              date(2028, 1, 1), '9'), (1, 1, 0))
        self.assertEqual(service.events.return_value.insert.call_count, 1)

    def test_duplicate_in_input_only_insert_once(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {'items': []}
        first = dict(summary='Regata', start=date(2027,5,2), end=date(2027,5,3),
                     description='',location='',international=False,source='https://veslo.cz/a')
        second = dict(first, summary=' regata ', end=date(2027,5,4), source='https://veslo.cz/b')
        self.assertEqual(sync(service,'calendar',[first,second],date(2027,1,1),date(2028,1,1),'9'), (1,1,0))
        self.assertEqual(service.events.return_value.insert.call_count, 1)

    def test_existing_timed_event_blocks_import(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {
            'items':[{'summary':'REGATA', 'start':{'dateTime':'2027-05-02T09:00:00+02:00'},
                      'end':{'dateTime':'2027-05-02T16:00:00+02:00'}}]}
        event = dict(summary='Regata',start=date(2027,5,2),end=date(2027,5,4),
                     description='',location='',international=False,source='https://veslo.cz/a')
        self.assertEqual(sync(service,'calendar',[event],date(2027,1,1),date(2028,1,1),'9'),(0,1,0))
        service.events.return_value.insert.assert_not_called()

    def test_duplicate_no_delete(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {'items':[{'summary':'Regata', 'start':{'date':'2026-11-03'}, 'end':{'date':'2026-11-04'}}]}
        event = dict(summary='Regata',start=date(2026,11,3),end=date(2026,11,4),description='',location='',international=False,source='https://veslo.cz/test')
        result = sync(service,'calendar',[event],date(2026,10,16),date(2027,10,16),'9')
        self.assertEqual(result,(0,1,0))
        service.events.return_value.delete.assert_not_called()
        service.events.return_value.insert.assert_not_called()

class CancellationTests(unittest.TestCase):
    def make_workbook(self, values, style=None):
        wb = Workbook()
        ws = wb.active
        ws.title = '2027'
        ws.append(['datum', 'místo', 'název akce', 'pořadatel', 'stav'])
        for day, title, state in values:
            ws.append([day, 'Praha', title, 'ČVS', state])
            if style:
                style(ws, ws.max_row)
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def events(self, data):
        return workbook_events(data, {2027}, date(2027, 1, 1), date(2028, 1, 1),
                               'https://www.veslo.cz/data/test.xlsx')

    def test_cancelled_status_in_title_or_note(self):
        result = self.events(self.make_workbook([
            ('5.6.', 'ZRUŠENO !!! Pražská regata ZRUŠENO !!!', None),
            ('6.6.', 'Regata juniorů', 'Zrušeno z důvodu počasí'),
            ('7.6.', 'Kontrolní závody seniorů A', None),
        ]))
        self.assertEqual([e['cancelled'] for e in result], [True, True, False])
        self.assertEqual(canonical_title(result[0]['summary']), canonical_title('Pražská regata'))

    def test_yellow_highlight_and_red_text(self):
        def style(ws, idx):
            if idx == 2:
                ws.cell(idx, 3).fill = PatternFill(fill_type='solid', fgColor='FFFFFF00')
            if idx == 3:
                ws.cell(idx, 3).font = Font(color='FFFF0000')
        items = self.events(self.make_workbook([
            ('10.6.', 'Regata žlutá', None), ('11.6.', 'Regata červená', None)], style))
        self.assertTrue(all(e['cancelled'] for e in items))

    def test_struck_or_red_fill(self):
        def style(ws, idx):
            if idx == 2:
                ws.cell(idx, 3).font = Font(strike=True)
            if idx == 3:
                ws.cell(idx, 1).fill = PatternFill(fill_type='solid', fgColor='FFFFC7CE')
        items = self.events(self.make_workbook([
            ('10.6.', 'Regata přeškrtnutá', None), ('11.6.', 'Regata růžově červená', None)], style))
        self.assertTrue(all(e['cancelled'] for e in items))

    def test_plain_and_blue_not_cancelled(self):
        def style(ws, idx):
            if idx == 2:
                ws.cell(idx, 3).font = Font(color='FF2F75B5')
        items = self.events(self.make_workbook([
            ('10.6.', 'Regata modrá', None), ('11.6.', 'Regata běžná', None)], style))
        self.assertFalse(any(e['cancelled'] for e in items))

    def test_theme_red_highlight(self):
        from gcalveslo import is_alert_color
        self.assertTrue(is_alert_color(Color(theme=5, tint=0.5), {5:'FF0000'}))
        self.assertTrue(is_alert_color(Color(rgb='00800000')))
        self.assertFalse(is_alert_color(Color(theme=5, tint=0.5), {5:'008000'}))

    def test_cancelled_source_beats_active_from_other_workbook(self):
        active = self.make_workbook([('5.6.', 'Regata v Praze', '')])
        cancelled = self.make_workbook([('5.6.', 'ZRUŠENO !!! Regata v Praze', '')])
        urls = ['https://www.veslo.cz/2027a.xlsx', 'https://www.veslo.cz/2027b.xlsx']
        def fetch(_, url):
            return dict(zip(urls, (active, cancelled)))[url]
        with patch('gcalveslo.sources', return_value=urls), patch('gcalveslo.get', side_effect=fetch):
            events, _ = load_events(Mock(), 'https://www.veslo.cz/source', date(2027, 1, 1))
        self.assertEqual(len(events), 1)
        self.assertTrue(events[0]['cancelled'])

    def test_cancelled_not_added_and_managed_old_event_deleted(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {'items': [
            {'id': 'old123', 'summary': 'Regata v Praze', 'start': {'date': '2027-06-05'},
             'extendedProperties': {'private': {'gcalveslo': '1'}}}]}
        cancelled = dict(summary='ZRUŠENO !!! Regata v Praze', start=date(2027,6,5),
                         end=date(2027,6,6), description='', location='',
                         international=False, source='https://www.veslo.cz/1',cancelled=True)
        self.assertEqual(sync(service,'mycal',[cancelled],date(2027,1,1),date(2028,1,1),'9'), (0,1,1))
        service.events.return_value.delete.assert_called_once_with(calendarId='mycal',
                                                                    eventId='old123', sendUpdates='none')
        service.events.return_value.insert.assert_not_called()

    def test_all_managed_duplicates_deleted_unmanaged_kept(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {'items': [
            {'id': 'first', 'summary': 'Regata', 'start': {'date':'2027-06-05'},
             'extendedProperties':{'private':{'gcalveslo':'1'}}},
            {'id': 'second', 'summary': 'Regata', 'start': {'date':'2027-06-05'},
             'description':'Pořadatel\nZdroj: https://www.veslo.cz/data/test.xlsx'},
            {'id': 'personal', 'summary': 'Regata', 'start':{'date':'2027-06-05'}}]}
        cancelled = dict(summary='Regata', start=date(2027,6,5), end=date(2027,6,6),
                         cancelled=True, description='',location='',international=False,
                         source='https://www.veslo.cz/data/test.xlsx')
        self.assertEqual(sync(service, 'mycal', [cancelled], date(2027,1,1), date(2028,1,1), '9'),(0,1,2))
        calls = service.events.return_value.delete.call_args_list
        self.assertEqual({c.kwargs['eventId'] for c in calls}, {'first','second'})

    def test_cancelled_wins_even_if_direct_sync_active_duplicate(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {'items': []}
        race = dict(summary='Regata',start=date(2027,6,5),end=date(2027,6,6),
                    description='',location='',international=False,source='https://veslo.cz/a')
        result = sync(service, 'mycal', [race, dict(race, cancelled=True)],
                      date(2027,1,1),date(2028,1,1),'9')
        self.assertEqual(result, (0,2,0))
        service.events.return_value.insert.assert_not_called()

    def test_dry_run_does_not_delete(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {'items': [
            {'id':'old','summary':'Regata','start':{'date':'2027-06-05'},
             'extendedProperties':{'private':{'gcalveslo':'1'}}}]}
        race = dict(summary='Regata',start=date(2027,6,5),end=date(2027,6,6),
                    description='',location='',international=False,source='https://veslo.cz/a',cancelled=True)
        self.assertEqual(sync(service,'mycal',[race],date(2027,1,1),date(2028,1,1),'9',dry_run=True),(0,1,1))
        service.events.return_value.delete.assert_not_called()

    def test_missing_source_race_is_not_deleted(self):
        service = Mock()
        service.events.return_value.list.return_value.execute.return_value = {'items': [
            {'id':'old','summary':'Regata','start':{'date':'2027-06-05'},
             'extendedProperties':{'private':{'gcalveslo':'1'}}}]}
        self.assertEqual(sync(service,'mycal',[],date(2027,1,1),date(2028,1,1),'9'),(0,0,0))
        service.events.return_value.delete.assert_not_called()


if __name__=='__main__': unittest.main()
